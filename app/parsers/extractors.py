import os
import io
import pymupdf
import docx
import openpyxl
import pptx
from PIL import Image
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.parsers.registry import BaseExtractor, ExtractedDocument, ExtractedNode, registry
# Import our robust PDF Parser to serve as the main engine for PDFs
from app.parsers.pdf_parser import LegalDocumentParser

class TextExtractor(BaseExtractor):
    def __init__(self):
        self.splitter = RecursiveCharacterTextSplitter(chunk_size=1500, chunk_overlap=150)

    def extract(self, file_bytes: bytes, filename: str) -> ExtractedDocument:
        # Decode raw bytes directly to string
        text = file_bytes.decode("utf-8", errors="replace")
        chunks = self.splitter.split_text(text)
        nodes = [ExtractedNode(text=chunk, metadata={"page_number": 1, "page_label": "1", "source": filename, "format": "text"}) for chunk in chunks]
        return ExtractedDocument(filename=filename, nodes=nodes)

class PDFExtractor(BaseExtractor):
    def __init__(self):
        self.robust_parser = LegalDocumentParser()

    def extract(self, file_bytes: bytes, filename: str):
        return self.robust_parser.parse_pdf(file_bytes, filename)

class ImageExtractor(BaseExtractor):
    def __init__(self):
        self.splitter = RecursiveCharacterTextSplitter(chunk_size=1500, chunk_overlap=150)

    def extract(self, file_bytes: bytes, filename: str) -> ExtractedDocument:
        Image.MAX_IMAGE_PIXELS = 100000000
        # 🔥 MEMORY OPTIMIZATION: Disabled local OCR for 512MB Free Tier. 
        # Cannot run pytesseract on images in this environment.
        text = "[Local OCR disabled on Render free tier. Cannot extract text from image.]"
        chunks = self.splitter.split_text(text)
        nodes = [ExtractedNode(text=chunk, metadata={"page_number": 1, "page_label": "Image", "source": filename, "format": "image"}) for chunk in chunks]
        return ExtractedDocument(filename=filename, nodes=nodes)

class DocxExtractor(BaseExtractor):
    def __init__(self):
        self.splitter = RecursiveCharacterTextSplitter(chunk_size=1500, chunk_overlap=150)

    def extract(self, file_bytes: bytes, filename: str) -> ExtractedDocument:
        # docx natively accepts BytesIO
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_data:
                    full_text.append(" | ".join(row_data))
        text = "\n".join(full_text)
        chunks = self.splitter.split_text(text)
        nodes = [ExtractedNode(text=chunk, metadata={"page_number": 1, "page_label": "Document", "source": filename, "format": "docx"}) for chunk in chunks]
        return ExtractedDocument(filename=filename, nodes=nodes)

class SpreadsheetExtractor(BaseExtractor):
    def __init__(self):
        self.splitter = RecursiveCharacterTextSplitter(chunk_size=1500, chunk_overlap=150)

    def extract(self, file_bytes: bytes, filename: str) -> ExtractedDocument:
        # openpyxl natively accepts BytesIO
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        nodes = []
        for i, sheet_name in enumerate(wb.sheetnames):
            sheet = wb[sheet_name]
            sheet_text = []
            for row in sheet.iter_rows(values_only=True):
                row_data = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                if row_data:
                    sheet_text.append(" | ".join(row_data))
            if sheet_text:
                text = "\n".join(sheet_text)
                chunks = self.splitter.split_text(text)
                for chunk in chunks:
                    nodes.append(ExtractedNode(text=chunk, metadata={"page_number": i+1, "page_label": f"Sheet: {sheet_name}", "source": filename, "format": "xlsx"}))
        return ExtractedDocument(filename=filename, nodes=nodes)

class PptxExtractor(BaseExtractor):
    def __init__(self):
        self.splitter = RecursiveCharacterTextSplitter(chunk_size=1500, chunk_overlap=150)

    def extract(self, file_bytes: bytes, filename: str) -> ExtractedDocument:
        # python-pptx natively accepts BytesIO
        prs = pptx.Presentation(io.BytesIO(file_bytes))
        nodes = []
        for i, slide in enumerate(prs.slides):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text.strip())
            if slide_text:
                text = "\n".join(slide_text)
                chunks = self.splitter.split_text(text)
                for chunk in chunks:
                    nodes.append(ExtractedNode(text=chunk, metadata={"page_number": i+1, "page_label": f"Slide {i+1}", "source": filename, "format": "pptx"}))
        return ExtractedDocument(filename=filename, nodes=nodes)

# Register Extractors Securely
registry.register("application/pdf", PDFExtractor)
registry.register("text/plain", TextExtractor)
registry.register("image/jpeg", ImageExtractor)
registry.register("image/png", ImageExtractor)
registry.register("image/tiff", ImageExtractor)
registry.register("image/webp", ImageExtractor)
registry.register("application/vnd.openxmlformats-officedocument.wordprocessingml.document", DocxExtractor)
registry.register("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", SpreadsheetExtractor)
registry.register("application/vnd.openxmlformats-officedocument.presentationml.presentation", PptxExtractor)