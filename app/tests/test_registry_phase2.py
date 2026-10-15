import os
import sys
import pytest
import docx
import openpyxl
import pptx

sys.path.append(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
from app.parsers.extractors import registry

def test_registry_docx_extraction(tmp_path):
    docx_file = tmp_path / "contract.docx"
    doc = docx.Document()
    doc.add_paragraph("This is a strict liability clause.")
    table = doc.add_table(rows=1, cols=2)
    row = table.rows[0].cells
    row[0].text = "Fee"
    row[1].text = "$5,000"
    doc.save(docx_file)
    
    extracted = registry.extract(str(docx_file), "contract.docx")
    assert len(extracted.nodes) > 0
    assert "strict liability clause" in extracted.nodes[0].text
    assert "Fee | $5,000" in extracted.nodes[0].text
    assert extracted.nodes[0].metadata["format"] == "docx"

def test_registry_xlsx_extraction(tmp_path):
    xlsx_file = tmp_path / "financials.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q3_Earnings"
    ws.append(["Revenue", "150000"])
    wb.save(xlsx_file)
    
    extracted = registry.extract(str(xlsx_file), "financials.xlsx")
    assert len(extracted.nodes) > 0
    assert "Revenue | 150000" in extracted.nodes[0].text
    assert extracted.nodes[0].metadata["page_label"] == "Sheet: Q3_Earnings"

def test_registry_pptx_extraction(tmp_path):
    pptx_file = tmp_path / "deck.pptx"
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Acquisition Strategy"
    prs.save(pptx_file)
    
    extracted = registry.extract(str(pptx_file), "deck.pptx")
    assert len(extracted.nodes) > 0
    assert "Acquisition Strategy" in extracted.nodes[0].text
    assert extracted.nodes[0].metadata["page_label"] == "Slide 1"
